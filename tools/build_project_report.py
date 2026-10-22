from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).resolve().parents[1]
doc=Document();section=doc.sections[0]
for style in doc.styles:
 for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.8)
section.left_margin=section.right_margin=Inches(.85)
for name in ['Normal','Body Text','List Bullet','List Number']:
 style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(11)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
 style.paragraph_format.line_spacing=1.2;style.paragraph_format.space_after=Pt(7)
for name,size in [('Title',26),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '17365D')
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');style.paragraph_format.keep_with_next=True
header=section.header.paragraphs[0];header.text='TOUCHLINE    触线足球经理';header.style='Caption'
footer=section.footer.paragraphs[0];footer.alignment=2;footer.add_run('项目报告  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text):return doc.add_paragraph(text)
def heading(text,level=1):return doc.add_heading(text,level)
def table(rows):
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 total=6.8
 widths=([1.55,total-1.55] if len(rows[0])==2 else [total/len(rows[0])]*len(rows[0]))
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,value in enumerate(row):
   cells[j].width=Inches(widths[j]);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   paragraph=cells[j].paragraphs[0];paragraph.paragraph_format.space_after=Pt(5);paragraph.paragraph_format.space_before=Pt(5)
   run=paragraph.add_run(value);run.font.size=Pt(10 if len(row)>5 else 10.5)
   tcpr=cells[j]._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if i==0 else 'F1F5F9' if i%2==0 else 'FFFFFF');tcpr.append(shade)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');margins.append(x)
   tcpr.append(margins)
   if i==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
  pr=t.rows[i]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if i==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
 t._tbl.tblPr.append(borders);doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc.add_paragraph('触线足球经理项目报告',style='Title')
p('项目总体介绍  智能体设计与配置  游戏操作说明')
p('开发验证版  |  2026年10月19日')
p('本报告面向项目展示、评审和首次使用者。触线足球经理是由原有MATLAB足球模拟规则迁移而来的Windows单机经营游戏，将可解释的比赛引擎、本地语言模型和俱乐部长期经营结合起来。当前已经可以创建世界、经营球队、观看比赛、保存进度并完成赛季流程，但仍属于开发验证版，不能据此宣称已经达到商业游戏的稳定性或长期平衡水平。')
heading('阅读导航',2)
p('第一部分介绍项目目标、规模与架构；第二部分说明本地智能体使用的模型、运行配置、能力、行为及保护规则；第三部分给出完整游戏说明，覆盖开局、数据、比赛、理念、战术、转会、年度结算和生涯模式。')
heading('一 项目总体介绍')
heading('1 项目定位与使用场景',2)
p('玩家扮演一家足球俱乐部的经理，围绕阵容、战术、资金、球员关系和赛季目标作出决策。其他俱乐部由本地智能体与自动教练共同经营。比赛采用抽象属性模拟和二维球场动画，不追求真实足球的逐帧物理运动，而强调经营决策、阵容搭配和不确定结果之间的联系。')
p('项目适合展示离线智能体在受约束环境中的感知、选择、工具调用和反馈过程。智能体不能直接改写存档或任意执行代码，只能提交预定义动作，由规则引擎审核并执行。')
table([['项目项','当前实现'],['平台与语言','64位Windows；C#与.NET Framework 4.x；WinForms图形界面'],['世界规模','30家俱乐部，3个联赛，新世界800名随机生成球员'],['赛事体系','联赛、一级杯赛、二级杯赛、欧冠、欧联及欧冠资格赛'],['比赛交互','三路攻防、球权和传球动画、分段推进、换人、战术调整'],['经营系统','转会、续约、工资、商业价值、荣誉、个人统计奖项及训练'],['生涯模式','长期培养；球星迭起，含年龄、成长、衰退、新秀与退役'],['保存方式','本地JSON封装存档、校验、自动保存及上一份备份'],['语言','中文和英文显示；自定义球队及球员名字保留原文']])
heading('2 模块架构与数据流',2)
p('界面首先接收玩家操作；赛程控制器确定当前赛事与阶段；智能体或自动教练输出候选动作；规则校验器检查阵容、权限、换人和资金约束；引擎执行比赛或经营结果；日志、统计和存档最后反映已经执行的状态。推进任务在副本中处理，成功后提交，失败或取消则保留上一次已提交状态。')
table([['模块','职责'],['Domain与Storage','球队、球员、赛事、交易与生涯状态；存读档和校验'],['Football与Competitions','属性加权、三路比赛、赛程、积分与淘汰晋级'],['SeasonFlow与CareerCycle','收入、奖项、好感、工资、成长衰退与新秀'],['Agents与AgentRequests','本地模型请求、结构化动作、合法性检查与重试'],['CoachingPolicy','分配模型或自动教练决策，并提供比赛观察信息'],['MainForm与MatchPresentation','经营界面、动画、换人、命名与交互'],['ClubChoices与Localization','方向抽签、战术商店、中英文显示资源']])
heading('3 设计特点与当前边界',2)
p('项目的主要特点是规则和智能体职责分离：比赛结果可追溯到属性与概率，智能体则负责在有限动作范围内表达经营偏好。隐藏性格、对称关系、不同阵容深度和年龄变化，使球队不只在当前战力上存在差异，也在预算和长期目标上形成区别。')
p('当前仍存在明确边界。模型规模较小，可能给出不合理的解约建议或与实际动作不一致的解释。自动排阵是属性估值工具，不是对最优胜率的证明。完整赛季数值回归不代表已完成多年经济平衡、全赛季智能体质量或商业发行兼容性验证。现有美术、声音、引导与无障碍体验也仍可继续完善。')
heading('二 所用智能体介绍')
heading('1 本地模型与运行配置',2)
table([['配置项','当前配置'],['模型','Qwen2.5 1.5B Instruct GGUF'],['量化文件','qwen2.5-1.5b-instruct-q4_k_m.gguf，Q4_K_M，1,117,320,736字节'],['推理程序','llama.cpp，项目安装记录版本b11719，Windows CPU x64'],['运行范围','本机回环地址127.0.0.1或localhost；默认端口17845'],['上下文与并发','8192 token上下文；单请求推理槽位'],['CPU线程','按处理器数量一半配置，最少2、最多8线程'],['生成参数','temperature为0.35；按决策场景设置150至500个输出token预算'],['响应约束','JSON结构化响应及枚举值限制；程序另做动作合法性验证'],['运行方式','安装阶段准备模型，游戏运行离线；内置服务不可用时检查并自动启动']])
p('上述信息来自项目代码及LocalAI安装清单，而非对模型官方能力的额外承诺。模型推理速度随CPU、内存和上下文长度变化；增加比赛观察信息会提高输入量，不能保证每次决策同样迅速。')
heading('2 感知信息与可执行行为',2)
p('智能体能读取本队资金、工资、阵容、属性、伤停、疲劳、好感、忠诚、理念、战术及近期经营记忆。比赛阶段提供对手与赛程资料，必要时补充当前比分、实际首发、可换入球员和三路能力估值。经营阶段提供可负担目标和待回应报价。模型只在这些输入及允许动作范围内选择，不会获得比赛未来随机结果。')
table([['阶段','可执行行为'],['赛前','从合法候选中选择阵容与战术；按升级点约束提出训练'],['中场','保持或更换战术；在剩余名额内提出换人'],['续约','保留或释放球员；球星迭起模式可提出主动退役'],['自主转会','处理球员主动提出的加盟方案'],['需求转会','挂牌、报价、接受或拒绝报价；遵守每窗5次有效主动报价上限']])
heading('3 性格 关系与记忆',2)
p('每家俱乐部拥有持久的隐藏经营性格，例如收益经营、竞技至上、欧冠雄心、联赛立足、忠诚传统、银河战舰、青训建设、稳健生存、杯赛猎手和体系建队。性格通过目标提示、候选采购排序和轮换方案影响决策。银河战舰仅分配给初始化时资金和商业价值达到要求的球队。性格并非硬性保证模型每次都作出同一种选择。')
p('任意两队共享对称关系值，越高越友好。同联赛初始范围为−99至100，跨联赛为1至200。关系供交易决策参考，不能突破工资和资金约束。当前关系不随比赛自动演化。近期经营记录提供有限记忆，但没有在游戏过程中训练模型权重或实现自主学习。')
heading('4 决策流程与保护措施',2)
p('一次模型决策依次经历读取当前状态、生成结构化请求、模型返回JSON、校验动作、执行与输出结果。校验包括本队球员归属、位置合法性、健康门将优先、战术是否解锁、最多5次换人、不能重复上场、报价次数，以及付款后能否负担全队工资。模型不能通过解释文字绕过这些规则。')
p('普通常规比赛使用自动教练减少等待，关键淘汰赛和经营阶段保留模型。对阵玩家的比赛始终由模型作赛前及中场选择，提供更完整比赛观察，但不会额外提高球员属性或修改比赛概率。此调度区别不会在游戏界面作为对手等级标签展示。')
p('连续三次非法动作后，允许回退的决定使用固定程序并记录原因；对阵玩家时则暂停等待重试，不以固定程序替代对手。请求失败与非法动作属于不同问题：内置服务未运行时自动检查并启动；连接错误显示具体端口和原因；超时和取消不会被伪装为已经执行的决定。')
heading('5 验证情况与后续方向',2)
p('现有测试覆盖首季588场、含欧战第二季734场、资格赛衔接、存读档、换人、转会、年龄结算、改名权限和战术商店；本地模型另有赛前、中场、续约和需求转会调用测试。原属性加权曾完成3000组对照。测试通过表示这些检查条件成立，不证明全部策略合理或所有极端状态安全。')
p('后续应重点评估多赛季经济走势、模型经营质量、特殊阵容下的决策、英文显示完整性和不同电脑上的推理延迟。可在保持规则校验的前提下比较更大模型或更好的候选工具，但不宜直接用模型生成比赛数值。')
doc.add_page_break();heading('三 游戏说明书')
manual=(root/'游戏说明书.md').read_text(encoding='utf-8-sig')
lines=manual.splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line or line.startswith('# '):continue
 if line.startswith('## '):
  title=re.sub(r'^\d+\.\s*','',line[3:]).replace('：',' ').replace('（',' ').replace('）','').replace('·',' ')
  heading(title,2);continue
 if line.startswith('|'):
  rows=[]
  while True:
   if not re.fullmatch(r'[| :\-]+',line):rows.append([v.strip().replace('`','') for v in line.strip('|').split('|')])
   if i>=len(lines) or not lines[i].strip().startswith('|'):break
   line=lines[i].strip();i+=1
  table(rows);continue
 if line.startswith('- '):doc.add_paragraph(line[2:].replace('`',''),style='List Bullet')
 else:p(line.replace('`',''))
heading('报告依据',2)
p('本报告依据当前项目src目录、LocalAI/manifest.json、游戏说明书及build/test-results验证记录整理。自定义世界数据仅用于本机游戏；原MATLAB、Excel和MAT文件不作为新程序运行时数据库，也不会由新程序自动回写。')
doc.save(root/'触线足球经理项目报告_最新版.docx')
print(root/'触线足球经理项目报告_最新版.docx')
