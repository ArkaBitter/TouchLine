from pathlib import Path
import sys
sys.stdout.reconfigure(encoding="utf-8")
from copy import deepcopy
import json,shutil
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
root=Path('FootballManager');out=root/'build/report-update-20261021';out.mkdir(parents=True,exist_ok=True)
path=root/'触线足球经理项目报告_最新版.docx';backup=out/'修改前项目报告.docx'
if not backup.exists():shutil.copy2(path,backup)
d=Document(backup);paras=list(d.paragraphs);changed=[]
def replace(p,text):
 old=p.text
 if p.runs:
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
 else:p.add_run(text)
 changed.append({'before':old,'after':text})
def setp(i,text):replace(paras[i],text)
def cell(t,r,c,text):
 p=d.tables[t].cell(r,c).paragraphs[0];replace(p,text)
 for extra in d.tables[t].cell(r,c).paragraphs[1:]:extra._p.getparent().remove(extra._p)
def insert_after(p,text):
 from docx.text.paragraph import Paragraph
 from docx.oxml import OxmlElement
 node=OxmlElement('w:p');p._p.addnext(node);new=Paragraph(node,p._parent);new.style=d.styles['Normal'];new.add_run(text);return new
setp(2,'开发验证版  |  2026年10月21日')
setp(3,'本报告面向项目展示、评审和首次使用者。触线足球经理由MATLAB足球模拟规则迁移而来，是结合比赛引擎、本地语言模型和俱乐部经营的Windows单机游戏。目前已实现创建世界、经营球队、观看比赛、保存进度和赛季结算功能；现行版本完整流程验收仍在进行，尚不代表商业游戏稳定性或长期平衡已获验证。')
setp(25,paras[25].text+'各队持有独立经营状态，共享同一个本地模型服务，而非分别加载29份模型。')
setp(12,'界面接收玩家操作，赛程控制器确定赛事与阶段；候选工具生成合法阵容、换人、换位、训练或交易选项，模型或自动教练选择动作。规则引擎再次校验后执行，输出比赛结果、公开摘要与存档。推进任务在副本中处理，每场已完成的电脑比赛另存检查点；后续失败或取消只撤销尚未提交的处理，不抹去已提交的比赛。')
setp(16,'当前仍有明确边界。小模型可能作出不理想的轮换、解约或经营选择；合法决定不等于高质量策略。公开摘要已由实际结构化动作生成，避免原始解释与执行内容不一致，但这不能证明模型判断正确。阵容权重不是最优胜率证明。现行版本完整端到端验收尚未完成，多赛季平衡、不同电脑兼容性、美术和新手引导仍需继续验证。')
setp(20,'配置依据项目代码及本地安装清单。兼容CUDA环境优先使用GPU；不可用或启动失败时改为CPU推理，仍使用同一本地模型。进入已配置模型的游戏界面时可后台预热；预热只加载模型，不提前决定比赛或交易。推理时间随硬件、输入长度、输出动作数量及缓存命中变化，不保证固定响应时间。')
setp(22,'智能体读取本队资金、工资、阵容、属性、伤停、疲劳、好感、忠诚、理念、战术、经营方针及最近4条经营记忆。比赛阶段提供对手、比分、当前阵容、出场分钟、剩余换人名额、合法替补与三路能力等信息；六周规划提供排名与未来赛程；转会阶段提供可负担目标、关系值和待回应报价。模型无法读取未来随机结果，也不直接生成比赛比分。')
setp(28,'一次决策依次经历读取当前状态、生成候选与JSON请求、模型选择动作编号、解码、合法性校验、执行和输出。赛前选择整套阵容与战术方案；临场选择换人及换位候选，训练选择预算内的方案编号。报价逐笔选择接受、拒绝或暂缓，不能同时接受和拒绝同一笔。归属、伤停、门将优先、战术权限、5人换人上限、5次报价上限及转入后工资预算仍由程序校验；临场组合先统一检验，防止部分执行后才失败。')
setp(29,'普通比赛由自动教练处理赛前和中场。每队智能体可为未来六周额外选最多3场重点比赛；对阵玩家、杯赛半决赛或决赛、欧冠或欧联八强，以及赛季末符合条件的直接排名争夺另行保障模型调用。排名争夺判定为联赛第16周起、双方积分差不超过3分且跨越第1、2、4、6或8名边界。阵容使用主力模板不代表该场必定调用模型。')
insert_after(paras[29],'玩家对手在赛前，以及10、20、30、40、50、60、70、80分钟由本地模型决策，另保留当前2分钟回合制的44分钟中场节点，90分钟不再调用。模型可保持现状、换人、调整战术或选择一组外场球员换位；换位不占换人名额。节点随存档保存，重载不重复执行，旧存档不补执行已经过去的节点。玩家对手不使用固定教练替代模型，也不获得额外属性或概率优势。')
insert_after(paras[28],'为减少重复计算，每队按已拥有战术缓存主力和二线阵容。联赛、欧战及资格赛、一级杯赛淘汰赛、二级杯赛半决赛及以上优先主力模板，其他杯赛优先二线模板；赛前根据伤停、明显疲劳和低出场机会进行微调，临时伤停不删除默认主力。球员归属、能力、好感或战术配置变化会使模板重新生成。')
setp(30,'连续3次非法输出后，允许回退的决定使用固定程序并标记来源和原因；对阵玩家的赛前、中场和临场决定则暂停等待重试，不能由固定教练接管。连接错误、HTTP错误、超时和取消不当作非法动作兜底；内置服务遇到连接故障可重新检查并重试一次，仍失败则暂停。公开摘要及俱乐部记忆依据实际动作生成，原始解释仅作为诊断信息，不能取代真实执行结果。')
setp(32,'旧版数值回归及3000组属性加权对照提供历史验证依据，但不代表现行智能体已完成整季验收。端到端记录完成首赛季588场、资格赛72场及第二赛季第9周350场后暂停，过程跨版本，T01仍未完成。当前临场版本通过23项专项断言及既有AI可靠性回归；真实本地模型单场完成9次临场与中场决定，未使用固定教练接管。这些检查不能证明长期策略质量或全部边界状态安全。')
insert_after(paras[32],'提速专项在RTX 4070 Laptop 8GB、i7-14650HX上完成100个计时决定，未发生非法纠正、兜底或执行失败。CPU基线平均18.84秒，GPU与输入输出精简组合平均1.25秒；协议与动作工作量可能不同，不能将差值视为纯硬件加速比。约1.04秒的缓存组重复使用相同快照，不代表真实变化赛程的平均延迟。2与4并发已做隔离比较，当前正式服务仍为单推理槽位，不能宣称29个模型同时运行。')
setp(33,'后续重点是完成现行版本端到端验收，评估多赛季经济与策略质量，并在参赛电脑上检查离线启动、推理延迟和界面引导。模型可能过早用尽换人名额或提出欠佳经营动作，需通过重复评估判断，而不能仅用一次合法输出或一次速度测试证明其战术水平。')
setp(38,'双击“启动游戏.cmd”。在“本地AI设置”配置并启动模型，然后选择“新建随机世界”，或通过“读取本地存档”继续已有世界。新建时选择足球理念及俱乐部方向，再随机签约。正常运行离线，兼容CUDA时优先GPU，否则CPU；请保留完整FootballManager文件夹，不只拷贝exe。')
setp(41,'第一赛季没有欧冠或欧联，赛季末举行资格赛，第二赛季起举办欧战。程序已移除内置原版世界入口及名单资源。原版世界仅为本机私人存档，不随项目分发；未指定执教球队的存档读取后先选择理念与方向，普通存档继续原执教球队。新程序不回写原Excel、MAT或MATLAB文件。')
setp(49,'本地AI设置：配置本地推理服务，查看启动和诊断日志。诊断日志有容量上限与轮转，记录启动、连接、超时及非法输出等原因。')
setp(50,'推进日志：区分本地智能体、自动教练和非法输出后的固定回退，展示实际阵容、战术、训练、经营动作及比赛结果。比赛右侧可查看对手已经执行的战术、换人和换位，记录随存档保存；公开摘要不等于模型内部推理过程。')
setp(51,'点击“推进赛程”推进当前阶段。玩家比赛按段推进，暂停时可换人、互换场上位置或打开“战术与效果”查看并应用已有战术；预览与取消不会改变战术。换位不消耗换人名额，每场最多换5人，换下不能重返赛场。右侧提供双方球员名字、位置、身价和能力。足球动画对应实际事件，射偏后球飞出球场再回到中圈。')
setp(57,'每周结束伤停减1、疲劳清零。每场结束按实际出场分钟结算疲劳：不超过30分钟增加0，超过30至60分钟增加1，超过60分钟增加3；累计疲劳降低本周后续比赛的四项正向能力，黑点不变，换位本身不增加疲劳。旧进行中存档按换人记录恢复分钟，无法可靠恢复时提示读取赛前存档。实际出场球员赛后有2.5%的随机伤病概率，持续2～5周；战术犯规另产生4周停赛。剩余伤停在商业收入阶段清除，之后资格赛仍可能产生新伤停。')
setp(91,'长期培养：不结算年龄、自然衰退、年度新秀和年龄成长，不开启退役功能；旧存档及私人原版世界存档默认此模式。比赛训练等原有成长机制照常保留。')
setp(104,'本地模型选择六周重点赛程和训练方案、关键比赛阵容、临场战术与换人换位，以及续约、退役、挂牌、报价和报价回应。AI训练仅在第6、12、18周比赛结束后的规划节点安排，可保留升级点；首个六周区间也有重点赛程规划，节点记录防止重复调用或扣点。玩家仍可随时使用升级点。非玩家决定连续3次非法可明确回退；玩家比赛相关决定暂停重试；连接故障或取消不伪装成正常决定。')
setp(113,'比赛决策分为模型选择与自动教练执行：每六周最多额外选3场重点比赛，并在规定节点分配训练；普通比赛减少模型调用，玩家比赛和指定关键比赛保留模型。赛前候选使用缓存阵容并按伤停、疲劳和出场机会微调。公开行为摘要来自结构化动作，不能作为模型真实思考过程的转录。')
setp(120,'已配置内置模型时，进入游戏界面可后台加载，推进前仍检查服务是否就绪；自定义本地服务需自行启动。兼容CUDA环境优先GPU，失败回到CPU。输入压缩为列名与数据行但保留全部事实；GPU路径可省去原始解释输出，CPU路径保留说明协议，公开动作摘要均保留。提示缓存只复用推理上下文，不复用过期的比赛或交易决定。')
setp(122,'本报告按2026年10月21日项目代码核对，主要依据CoachingPolicy、StrategicPlanning、LiveCoaching、SquadTemplates、Agents、InferenceTuning、MatchMinutes及DecisionNarration。验证依据为build/test-results/live-coaching、inference-20261021、private-world及待完成任务.txt；未为本次文档更新重新运行完整赛季或模型性能基准。原版私人存档不作为发布素材，报告不包含原版球员或俱乐部名单。')
cell(0,4,1,'三路攻防、球权和传球动画、分段推进、5人换人、场上换位与战术调整')
cell(1,5,1,'按赛事与重点赛程分配模型或自动教练；提供对手、阵容及比赛观察')
cell(2,3,1,'llama.cpp b11719；兼容CUDA时优先GPU，不可用时CPU')
cell(2,5,1,'8192 token上下文；正式服务单推理槽位，非29个并行模型')
cell(2,7,1,'temperature 0.35；普通赛前150 token，玩家比赛500，其他按场景预算')
cell(2,9,1,'离线运行；后台预热、就绪检查、GPU失败转CPU及提示缓存')
cell(3,1,1,'从合法的主力或二线模板候选中选择整套阵容与战术；此阶段不分配训练点')
cell(3,2,0,'中场与临场')
cell(3,2,1,'保持或调整战术，选择预算内换人及一组外场换位；玩家对手每10分钟另作决定')
t=d.tables[3];row=deepcopy(t.rows[1]._tr);t.rows[1]._tr.addprevious(row)
cell(3,1,0,'六周规划');cell(3,1,1,'未来六周额外重点比赛最多3场；第6、12、18周末选择合法训练方案，可保留升级点')
cell(4,10,1,'按场次实际分钟累计：≤30分钟加0，>30至60分钟加1，>60分钟加3；每周清零，影响后续比赛')
# Prevent an isolated table header at a page bottom and restore alternating rows.
for table in d.tables:
 for c in table.rows[0].cells:
  for p in c.paragraphs:p.paragraph_format.keep_with_next=True
for i,row in enumerate(d.tables[3].rows):
 if not i:continue
 for c in row.cells:
  shade=c._tc.get_or_add_tcPr().find(qn('w:shd'))
  if shade is not None:shade.set(qn('w:fill'),'F1F5F9' if i%2==0 else 'FFFFFF')
for section in d.sections:
 for p in section.header.paragraphs:
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
# Keep source layout and table treatments; normalize heading colors required for report output.
for style in d.styles:
 if style.name=='Title' or style.name.startswith('Heading'):
  style.font.color.rgb=RGBColor(0,0,0)
for p in d.paragraphs:
 if p.style.name=='Title' or p.style.name.startswith('Heading'):
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for index,note in [(20,'依据Agents中的LocalRuntime及InferenceTuning更新GPU回退、预热和缓存。'),(29,'依据CoachingPolicy与LiveCoaching更新模型调用范围；注意阵容模板级别与模型调用策略是两套规则。'),(32,'纠正完整赛季验收表述；只引用已有专项记录，T01仍未完成。'),(41,'原版快照已移为项目之外的私人存档，发布包不再含名单。'),(57,'依据MatchMinutes.Fatigue更新分钟阈值，避免沿用所有出场者统一加3的旧规则。')]:
 if hasattr(d,'add_comment'):d.add_comment(paras[index].runs[0],text=note,author='项目维护',initials='TL')
d.save(path)
(out/'changes.json').write_text(json.dumps({'ParagraphAndCellEdits':len(changed),'Changes':changed},ensure_ascii=False,indent=2),encoding='utf-8')
print('Updated paragraphs/cells:',len(changed));print('Document:',str(path))
